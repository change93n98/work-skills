"""Read back the actual XLSX exporter, including WAIT and overlap accounting."""
import json
import os
from pathlib import Path
import subprocess
import tempfile
import unittest

from kernel_categories import KernelClassifier
from prof import classify_rows,window_coverage

RUNTIME=Path.home()/'.cache/codex-runtimes/codex-primary-runtime/dependencies/node'
NODE=RUNTIME/('bin/node.exe' if os.name=='nt' else 'bin/node')
MODULES=RUNTIME/'node_modules'


def fixture_report():
    classifier=KernelClassifier(platform='teco')
    raw=[dict(name='teco_slave_gemm(Args)',record_count=1,total_duration_us=40,
              avg_duration_us=40,kernel_relative_pct=66.67,window_absolute_pct=40),
         dict(name='WAIT',record_count=1,total_duration_us=20,
              avg_duration_us=20,kernel_relative_pct=33.33,window_absolute_pct=20)]
    kernels,summary,excluded,busy=classify_rows(raw,classifier)
    coverage=window_coverage([dict(ts=0,dur=40000),dict(ts=30000,dur=20000,name='WAIT')],0,100000)
    return dict(mode='window',source='fixture.trace.json',metadata=dict(stage='decode',start_us='0',end_us='100',
                window_duration_us=100,window_coverage=coverage),kernels=kernels,summary=summary,excluded=excluded,
                busy_total_duration_us=busy,raw_total_duration_us=60,classification_rules=classifier.get_rule_provenance())


def moe_fixture_report():
    report=fixture_report();classifier=KernelClassifier(platform='teco')
    raw=[dict(name=name,record_count=1,total_duration_us=dur,avg_duration_us=dur,
              kernel_relative_pct=dur/80*100,window_absolute_pct=dur) for name,dur in [
        ('teco_slave_gemm(Args)',40),
        ('void tecolmk_slave_gemm_grouped_batched_v2_weight_half<3, 64>(GemmGroupBatchArgsV2<64>)',20),
        ('WAIT',20)]]
    kernels,summary,excluded,busy=classify_rows(raw,classifier)
    report.update(kernels=kernels,summary=summary,excluded=excluded,busy_total_duration_us=busy,raw_total_duration_us=80)
    report['metadata']['window_coverage']=window_coverage([
        dict(ts=0,dur=40000),dict(ts=30000,dur=20000),dict(ts=40000,dur=20000)],0,100000)
    return report

@unittest.skipUnless(NODE.is_file() and (MODULES/'@oai/artifact-tool').is_dir(),
                     'artifact-tool bundled runtime not available')
class TimeAccountingExportTests(unittest.TestCase):
    def export_report(self,report,root,stem='result'):
        source=root/(stem+'.json');output=root/(stem+'.xlsx')
        source.write_text(json.dumps(report,ensure_ascii=False),encoding='utf8')
        result=subprocess.run([str(NODE),str(Path(__file__).with_name('export_prof_report.mjs')),
                              str(source),str(output),str(MODULES)],capture_output=True,encoding='utf8')
        self.assertEqual(result.returncode,0,result.stderr[-2000:])
        return output

    def test_wait_retained_and_every_time_account_matches_source(self):
        import openpyxl
        report=fixture_report()
        with tempfile.TemporaryDirectory() as directory:
            root=Path(directory);source=root/'input.json';output=root/'result.xlsx'
            source.write_text(json.dumps(report,ensure_ascii=False),encoding='utf8')
            result=subprocess.run([str(NODE),str(Path(__file__).with_name('export_prof_report.mjs')),
                    str(source),str(output),str(MODULES)],capture_output=True,encoding='utf8')
            self.assertEqual(result.returncode,0,result.stderr[-2000:])
            wb=openpyxl.load_workbook(output,data_only=True)
            try:
                sheet=wb['分类汇总']
                for cell,expected in [('F12',40),('F13',20),('F14',60),('F15',-10),
                                      ('F16',50),('F17',50),('F18',100),('I17',.5),('I18',1)]:
                    self.assertAlmostEqual(sheet[cell].value,expected,msg=cell)
                self.assertEqual(sheet['C13'].value,'等待/同步 WAIT/SET')
                self.assertEqual(sheet['C17'].value,'空泡')
                self.assertEqual(wb.sheetnames,['分类汇总','Kernel明细','选窗依据'])
                self.assertEqual(sheet.freeze_panes,'A6')
                self.assertEqual(wb['Kernel明细'].freeze_panes,'A2')
                self.assertTrue(all(s.sheet_view.showGridLines is False for s in wb))
                self.assertEqual(sheet['A5'].font.name,'Arial')
                self.assertEqual(sheet['A5'].font.sz,10)
                self.assertEqual(sheet['A5'].font.color.rgb,'FFFFFFFF')
                self.assertEqual(sheet['A5'].fill.fgColor.rgb,'FF253B58')
                self.assertEqual(sheet['A2'].font.sz,14)
                self.assertTrue(sheet['G17'].font.bold)
                self.assertEqual(sheet['G6'].number_format,'#,##0.000')
                self.assertEqual(sheet['I17'].number_format,'0.00%')
                self.assertEqual(sheet.tables['CategorySummary'].tableStyleInfo.name,'TableStyleMedium2')
                self.assertTrue(sheet.tables['CategorySummary'].tableStyleInfo.showRowStripes)
                widths=[9,32,26,15,15,23,23,20,18,20,15]
                for i,width in enumerate(widths,1):
                    self.assertEqual(sheet.column_dimensions[openpyxl.utils.get_column_letter(i)].width,width)
                self.assertEqual(wb['选窗依据']['B9'].value,'llm-prof-classification-v2')
            finally:wb.close()

    def test_multiple_ranks_stay_separate_and_mixed_windows_are_rejected(self):
        import openpyxl
        report=fixture_report()
        with tempfile.TemporaryDirectory() as directory:
            root=Path(directory)
            output=self.export_report(dict(mode='batch',reports=[dict(report,rank=0,label='decode-step2-device'),
                                          dict(report,rank=1,label='decode-step2-device')]),root)
            wb=openpyxl.load_workbook(output,data_only=True)
            try:
                self.assertEqual(wb['分类汇总']['F6'].value,40)
                self.assertEqual(wb['分类汇总']['F19'].value,40)
                self.assertEqual(wb['分类汇总']['F18'].value,100)
                self.assertEqual(wb['分类汇总']['F31'].value,100)
                self.assertEqual(wb['Kernel明细'].max_row,5)
            finally:wb.close()
            bad=root/'bad.json';bad_output=root/'bad.xlsx'
            bad.write_text(json.dumps(dict(mode='batch',reports=[dict(report,rank=0,label='decode-step2'),
                            dict(report,rank=1,label='decode-step3')])),encoding='utf8')
            result=subprocess.run([str(NODE),str(Path(__file__).with_name('export_prof_report.mjs')),
                    str(bad),str(bad_output),str(MODULES)],capture_output=True,encoding='utf8')
            self.assertNotEqual(result.returncode,0)
            self.assertFalse(bad_output.exists())

    def test_imported_xlsx_preserves_source_and_marks_unknown_bubble(self):
        import openpyxl
        from copy import copy
        from prof import read_xlsx
        classifier=KernelClassifier(platform='teco')
        with tempfile.TemporaryDirectory() as directory:
            root=Path(directory)
            original=self.export_report(fixture_report(),root,'original')
            raw,metadata=read_xlsx(original,'Kernel明细')
            kernels,summary,excluded,busy=classify_rows(raw,classifier)
            output=self.export_report(dict(mode='xlsx',source=str(original),metadata=metadata,kernels=kernels,
                summary=summary,excluded=excluded,busy_total_duration_us=busy,
                classification_rules=classifier.get_rule_provenance()),root,'reclassified')
            before=openpyxl.load_workbook(original);after=openpyxl.load_workbook(output)
            try:
                self.assertEqual(after.sheetnames[:3],before.sheetnames)
                for sheet in before:
                    for row in sheet:
                        for cell in row:
                            actual=after[sheet.title][cell.coordinate]
                            for attr in ('value','font','fill','border','alignment','number_format'):
                                self.assertEqual(copy(getattr(cell,attr)),copy(getattr(actual,attr)),
                                                 (sheet.title,cell.coordinate,attr))
                self.assertEqual(after['Kernel明细']['P1'].value,'最新分类结果')
                self.assertEqual(after['最新分类汇总']['A17'].value,'空泡')
                self.assertEqual(after['最新分类汇总']['C17'].value,'n.a.')
                self.assertEqual(after['最新分类汇总']['A5'].fill.fgColor.rgb,'FF253B58')
            finally:before.close();after.close()

    def test_moe_group_row_denominator_and_overlap_accounting(self):
        import openpyxl
        report=moe_fixture_report()
        with tempfile.TemporaryDirectory() as directory:
            output=self.export_report(report,Path(directory))
            w=openpyxl.load_workbook(output,data_only=True)
            try:
                s=w['分类汇总'];self.assertEqual(s['C7'].value,'group gemm')
                self.assertEqual(s['C13'].value,'七类合计');self.assertEqual(s['H5'].value,'七类内占比')
                for cell,value in [('F6',40),('F7',20),('F13',60),('F14',20),('F15',80),
                                   ('F16',-20),('F17',60),('F18',40),('F19',100),('H7',1/3),('I7',.2)]:
                    self.assertAlmostEqual(s[cell].value,value,msg=cell)
                self.assertEqual(s['A5'].fill.fgColor.rgb,'FF253B58');self.assertEqual(s.freeze_panes,'A6')
                self.assertEqual(s.tables['CategorySummary'].ref,'A5:K19')
            finally:w.close()

    def test_moe_multirank_blocks_shift_by_fourteen_rows(self):
        import openpyxl
        report=moe_fixture_report()
        batch=dict(mode='batch',reports=[dict(report,rank=0,label='decode-step2-device'),dict(report,rank=1,label='decode-step2-device')])
        with tempfile.TemporaryDirectory() as directory:
            w=openpyxl.load_workbook(self.export_report(batch,Path(directory)),data_only=True)
            try:
                s=w['分类汇总']
                for cell,value in [('F6',40),('F7',20),('F19',100),('F20',40),('F21',20),('F32',40),('F33',100)]:
                    self.assertAlmostEqual(s[cell].value,value,msg=cell)
                self.assertEqual(s['C21'].value,'group gemm');self.assertEqual(w['Kernel明细'].max_row,7)
            finally:w.close()

    def test_old_six_category_xlsx_gains_group_row_without_changing_old_data(self):
        import openpyxl
        from kernel_categories import DEFAULT_RULES_PATH,DENSE_CATEGORIES
        from prof import read_xlsx,RAW_COLUMNS
        with tempfile.TemporaryDirectory() as directory:
            root=Path(directory);report=moe_fixture_report()
            artifact=json.loads(DEFAULT_RULES_PATH.read_text(encoding='utf-8'))
            artifact.update(schema_version=1,categories=list(DENSE_CATEGORIES));artifact['compatibility_keywords'].pop('group gemm')
            for group,rules in artifact['platforms'].items():artifact['platforms'][group]=[r for r in rules if r['category']!='group gemm']
            rulefile=root/'old-rules.json';rulefile.write_text(json.dumps(artifact),encoding='utf-8')
            legacy=KernelClassifier(rulefile,platform='teco')
            raw=[{key:k[key] for key in RAW_COLUMNS} for k in report['kernels']]
            report['kernels'],report['summary'],report['excluded'],report['busy_total_duration_us']=classify_rows(raw,legacy)
            report['classification_rules']=legacy.get_rule_provenance()
            original=self.export_report(report,root,'original');original_bytes=original.read_bytes()
            raw,metadata=read_xlsx(original,'Kernel明细');new=KernelClassifier(platform='teco')
            kernels,summary,excluded,busy=classify_rows(raw,new)
            updated=self.export_report(dict(mode='xlsx',source=str(original),metadata=metadata,kernels=kernels,
                summary=summary,excluded=excluded,busy_total_duration_us=busy,classification_rules=new.get_rule_provenance()),root,'updated')
            self.assertEqual(original.read_bytes(),original_bytes)
            w=openpyxl.load_workbook(updated,data_only=True)
            try:
                self.assertEqual(w['Kernel明细']['G3'].value,'gemm');self.assertEqual(w['Kernel明细']['P3'].value,'group gemm')
                self.assertEqual(w['分类汇总']['F6'].value,60)
                s=w['最新分类汇总'];self.assertEqual(s['A7'].value,'group gemm');self.assertEqual(s['C7'].value,20)
                self.assertEqual(s['A13'].value,'七类合计');self.assertEqual(s['C13'].value,60)
                self.assertEqual(s['A18'].value,'空泡');self.assertEqual(s['C18'].value,'n.a.')
            finally:w.close()


if __name__=='__main__':unittest.main()

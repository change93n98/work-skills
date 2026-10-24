#!/usr/bin/env python3
"""Explicit Perfetto window or exported-XLSX classification, without the UI."""
import argparse
import csv
import json
import math
import os
from pathlib import Path
import re
import shutil
import subprocess
from decimal import Decimal, InvalidOperation

from kernel_categories import CATEGORIES, KernelClassifier, add_classification_args
from prof_analyze_perfetto import run_perfetto_query

RAW_COLUMNS = ('name', 'record_count', 'total_duration_us', 'avg_duration_us',
               'kernel_relative_pct', 'window_absolute_pct')
DEFAULT_CATEGORIES = ('kernel', 'gpu_memcpy', 'gpu_memset')


def us_to_ns(text):
    """Parse decimal microseconds without passing epoch or offsets through float."""
    try:
        value = Decimal(str(text)) * 1000
        if not value.is_finite() or value < 0 or value != value.to_integral_value():
            raise ValueError('time must be nonnegative, finite and precise to whole ns')
        return int(value)
    except (InvalidOperation, TypeError) as exc:
        raise ValueError(f'invalid microseconds: {text}') from exc


def resolve_trace(path):
    path = Path(path).expanduser().resolve()
    if path.is_dir():
        matches = sorted(p for p in path.iterdir() if p.is_file() and
                         p.name.lower().endswith(('.json', '.json.gz', '.pftrace', '.perfetto-trace')))
        if len(matches) != 1:
            raise ValueError(f'{path}: found {len(matches)} traces; specify one file (do not merge ranks)')
        path = matches[0]
    if not path.is_file():
        raise ValueError(f'trace does not exist: {path}')
    return path


def window_queries(start_ns, end_ns, categories=DEFAULT_CATEGORIES, track_ids=()):
    if not 0 <= start_ns < end_ns:
        raise ValueError('require 0 <= start < end')
    if not categories or any('runtime' in c.lower() or c.lower() in
            ('cpu','cpu_op','user_annotation','python_function') for c in categories):
        raise ValueError('only explicit device event categories are allowed, not CPU/runtime categories')
    literals = ', '.join("'" + str(c).replace("'", "''") + "'" for c in categories)
    tracks = ' AND track_id IN (' + ','.join(str(int(t)) for t in track_ids) + ')' if track_ids else ''
    cte = f'''WITH time_window AS (
  SELECT {start_ns} AS start_offset_ns, {end_ns} AS end_offset_ns
)
'''
    where = f'''FROM slice, trace_bounds, time_window tw
WHERE category IN ({literals})
  AND ts >= trace_bounds.start_ts + tw.start_offset_ns
  AND ts < trace_bounds.start_ts + tw.end_offset_ns{tracks}'''
    grouped = cte + '''SELECT name, COUNT(*) AS record_count,
  SUM(dur) / 1000.0 AS total_duration_us,
  SUM(dur) / COUNT(*) / 1000.0 AS avg_duration_us,
  ROUND(SUM(dur) * 100.0 / SUM(SUM(dur)) OVER (), 2) AS kernel_relative_pct,
  ROUND(SUM(dur) * 100.0 / (tw.end_offset_ns - tw.start_offset_ns), 2) AS window_absolute_pct
''' + where + '\nGROUP BY name\nORDER BY total_duration_us DESC;\n'
    events = cte + 'SELECT name, ts, dur, category, track_id\n' + where + '\nORDER BY ts;\n'
    return grouped, events


def analyze_window(tp, start_us, end_us, categories=DEFAULT_CATEGORIES, track_ids=()):
    start, end = us_to_ns(start_us), us_to_ns(end_us)
    sql, events_sql = window_queries(start, end, categories, track_ids)
    bounds = run_perfetto_query(tp, 'SELECT start_ts, end_ts FROM trace_bounds')[0]
    if end > bounds['end_ts'] - bounds['start_ts']:
        raise ValueError('requested window extends beyond trace_bounds; check relative time and units')
    rows = run_perfetto_query(tp, sql)
    events = run_perfetto_query(tp, events_sql)
    if not rows:
        raise ValueError('no device events in this window; check time origin, category and track selection')
    if any(e['dur'] < 0 for e in events):
        raise ValueError('selected events contain unfinished/negative durations; inspect the trace')
    metadata = dict(time_origin='trace_bounds.start_ts', start_us=str(start_us), end_us=str(end_us),
                    start_offset_ns=start, end_offset_ns=end, window_duration_us=(end-start)/1000,
                    trace_start_ts_ns=str(bounds['start_ts']), trace_end_ts_ns=str(bounds['end_ts']),
                    absolute_start_ts_ns=str(bounds['start_ts']+start),
                    absolute_end_ts_ns=str(bounds['start_ts']+end),
                    selection='start-inclusive/end-exclusive; full duration, no clipping',
                    device_categories=list(categories), track_ids=sorted({e['track_id'] for e in events}),
                    selected_event_count=len(events),
                    events_crossing_end=sum(e['ts']+e['dur'] > bounds['start_ts']+end for e in events),
                    selected_total_duration_us=sum(e['dur'] for e in events)/1000)
    return rows, metadata, sql


def read_xlsx(path, sheet_name=None):
    import openpyxl  # Reading only; artifact-tool preserves and authors the output.
    wb = openpyxl.load_workbook(path, read_only=True, data_only=True)
    try:
        sheets = [wb[sheet_name]] if sheet_name else wb.worksheets
        candidates = []
        for sheet in sheets:
            headers = list(next(sheet.iter_rows(values_only=True), ()))
            if 'name' in headers and 'total_duration_us' in headers and 'record_count' in headers:
                candidates.append((sheet, headers))
        if len(candidates) != 1:
            raise ValueError('specify --sheet: require exactly one sheet with name/record_count/total_duration_us headers in row 1')
        sheet, headers = candidates[0]
        rows = []
        source_columns, source_row_count = len(headers), 1
        for source_row, values in enumerate(sheet.iter_rows(min_row=2, values_only=True), 2):
            source_columns = max(source_columns,len(values))
            source_row_count = source_row
            if all(v is None for v in values):
                continue
            data = dict(zip(headers, values))
            if not isinstance(data.get('name'), str) or not data['name']:
                raise ValueError(f'missing kernel name at row {source_row}')
            row = {key:data.get(key) for key in RAW_COLUMNS}
            row['source_row'] = source_row
            rows.append(row)
        return rows, dict(source_sheet=sheet.title, source_columns=source_columns,
                          source_row_count=source_row_count, source_sheet_names=wb.sheetnames, headers=headers)
    finally:
        wb.close()


def classify_rows(rows, classifier):
    """Preserve raw SQL denominators; six-category relative % excludes runtime markers."""
    if not rows:
        raise ValueError('no kernel rows')
    summary = {c:dict(category=c, kernel_kinds=0, record_count=0, total_duration_us=0.0,
                      window_absolute_pct=0.0) for c in CATEGORIES}
    excluded = dict(category='非计算/排除', kernel_kinds=0, record_count=0,
                    total_duration_us=0.0, window_absolute_pct=0.0)
    result = []
    all_have_window = all(r.get('window_absolute_pct') is not None for r in rows)
    for raw in rows:
        dur, count = raw['total_duration_us'], raw['record_count']
        if (not isinstance(dur, (int,float)) or not math.isfinite(dur) or dur < 0 or
                not isinstance(count, (int,float)) or not math.isfinite(count) or count <= 0 or int(count) != count):
            raise ValueError(f"invalid duration/count for {raw['name']}")
        detail = classifier.classify_detail(raw['name'])
        label = detail['category'] or ('非计算/'+raw['name'] if raw['name'] in ('WAIT','SET') else '非计算/排除')
        row = dict(raw, category=label, classification=detail)
        result.append(row)
        bucket = summary[detail['category']] if detail['category'] else excluded
        bucket['kernel_kinds'] += 1
        bucket['record_count'] += int(count)
        bucket['total_duration_us'] += dur
        if all_have_window:
            value = raw['window_absolute_pct']
            if not isinstance(value, (int,float)) or not math.isfinite(value) or value < 0:
                raise ValueError('invalid window_absolute_pct')
            bucket['window_absolute_pct'] += value
    busy = sum(s['total_duration_us'] for s in summary.values())
    for s in summary.values():
        s['kernel_relative_pct'] = s['total_duration_us'] / busy * 100 if busy else None
        if not all_have_window:
            s['window_absolute_pct'] = None
    excluded['kernel_relative_pct'] = None
    if not all_have_window:
        excluded['window_absolute_pct'] = None
    total = sum(r['total_duration_us'] for r in result)
    assert math.isclose(total, busy + excluded['total_duration_us'], abs_tol=1e-8)
    return result, list(summary.values()), excluded, busy


def exporter_runtime(args):
    base = Path.home()/'.cache/codex-runtimes/codex-primary-runtime/dependencies/node'
    node = args.node or str(base/('bin/node.exe' if os.name=='nt' else 'bin/node'))
    modules = args.node_modules or str(base/'node_modules')
    if not Path(node).is_file():
        node = shutil.which('node')
    if not node or not Path(modules).is_dir():
        raise ValueError('XLSX needs artifact-tool: supply --node and --node-modules, or use --no-xlsx for JSON/CSV')
    return node, modules


def main(argv=None):
    parser = argparse.ArgumentParser(description=__doc__)
    modes = parser.add_subparsers(dest='mode', required=True)
    window = modes.add_parser('window', help='Perfetto UI-relative microsecond window, no manual export')
    window.add_argument('--trace', required=True, type=Path)
    window.add_argument('--start-us', required=True)
    window.add_argument('--end-us', required=True)
    window.add_argument('--stage', required=True, choices=('prefill','decode','all'))
    window.add_argument('--device-category', action='append', default=[])
    window.add_argument('--track-id', action='append', type=int, default=[])
    window.add_argument('--trace-processor', help='Optional native trace_processor_shell path')
    table = modes.add_parser('xlsx', help='Classify an exported XLSX, preserving original columns and file')
    table.add_argument('--xlsx', required=True, type=Path)
    table.add_argument('--sheet')
    for sub in (window,table):
        sub.add_argument('--output-dir', required=True, type=Path)
        sub.add_argument('--label', default='prof-report', help='Output filename prefix (letters, digits, dot, dash, underscore)')
        sub.add_argument('--no-xlsx', action='store_true')
        sub.add_argument('--node')
        sub.add_argument('--node-modules')
        sub.add_argument('--overwrite', action='store_true', help='Replace prior report outputs; never the source XLSX')
        add_classification_args(sub)
    args = parser.parse_args(argv)
    try:
        if not re.fullmatch(r'[A-Za-z0-9][A-Za-z0-9_.-]*',args.label):
            raise ValueError('unsafe --label')
        runtime = None if args.no_xlsx else exporter_runtime(args)
        out = args.output_dir.expanduser().resolve()
        outputs = [out/(args.label+suffix) for suffix in ('.json','.csv','.txt','.xlsx','.sql')]
        if not args.overwrite and any(p.exists() for p in outputs):
            raise ValueError('report exists: choose another label/output directory or --overwrite')
        classifier = KernelClassifier(args.kernel_rules,platform=args.platform)
        sql = None
        if args.mode == 'window':
            source = resolve_trace(args.trace)
            from perfetto.trace_processor import TraceProcessor, TraceProcessorConfig
            tp = TraceProcessor(trace=str(source), config=TraceProcessorConfig(
                bin_path=args.trace_processor,load_timeout=60))
            try:
                raw, metadata, sql = analyze_window(tp,args.start_us,args.end_us,
                    tuple(dict.fromkeys(DEFAULT_CATEGORIES+tuple(args.device_category))),args.track_id)
            finally:
                tp.close()
            metadata['stage'] = args.stage
        else:
            source = args.xlsx.expanduser().resolve()
            if outputs[3] == source:
                raise ValueError('output must differ from source XLSX')
            raw, metadata = read_xlsx(source,args.sheet)
        kernels, summary, excluded, busy = classify_rows(raw,classifier)
        if args.mode=='window':
            # Use full precision durations/window for category percentages, not rounded raw % sums.
            for item in summary+[excluded]:
                item['window_absolute_pct']=item['total_duration_us']/metadata['window_duration_us']*100
        report = dict(mode=args.mode, source=str(source), metadata=metadata,
            classification_rules=classifier.get_rule_provenance(), kernels=kernels,
            summary=summary, excluded=excluded, busy_total_duration_us=busy,
            raw_total_duration_us=sum(r['total_duration_us'] for r in kernels),
            notes=['Raw first six columns preserve the supplied SQL selection and denominator, including WAIT/SET.',
                   'Category relative percentage excludes runtime markers. Full duration is used even across window end.',
                   'Summed durations may overlap across streams/devices; window percentage is not utilization or bubble.',
                   'Classification is name-based; evidence/confidence/ambiguity are preserved, not invocation source attribution.'])
        out.mkdir(parents=True,exist_ok=True)
        report_path=out/(args.label+'.json')
        report_path.write_text(json.dumps(report,ensure_ascii=False,indent=2),encoding='utf8')
        with (out/(args.label+'.csv')).open('w',encoding='utf-8-sig',newline='') as f:
            writer=csv.DictWriter(f,fieldnames=RAW_COLUMNS+('category','rule_id','confidence','ambiguity'))
            writer.writeheader()
            for r in kernels:
                writer.writerow({**{k:r.get(k) for k in RAW_COLUMNS},'category':r['category'],
                    'rule_id':r['classification']['rule_id'],'confidence':r['classification']['confidence'],
                    'ambiguity':'; '.join(r['classification']['ambiguity_reasons'])})
        text=[f"{args.label}: {len(kernels)} kernel names, {sum(r['record_count'] for r in kernels)} calls",
              f"Window: {metadata.get('window_duration_us','unavailable')} us; stage={metadata.get('stage','supplied table')}",
              f"Device kernel sum excluding runtime markers: {busy/1000:.6f} ms"]
        for s in summary:
            pct=s['kernel_relative_pct']
            text.append(f"{s['category']}: {s['total_duration_us']/1000:.6f} ms, {pct if pct is not None else 'n.a.'}%")
        text.extend([f"Excluded: {excluded['total_duration_us']/1000:.6f} ms",*report['notes']])
        (out/(args.label+'.txt')).write_text('\n'.join(text)+'\n',encoding='utf8')
        if sql:
            (out/(args.label+'.sql')).write_text(sql,encoding='utf8')
        if runtime:
            node, modules=runtime
            command=[node,str(Path(__file__).with_name('export_prof_report.mjs')),str(report_path),
                     str(outputs[3]),modules]
            subprocess.run(command,check=True)
            if not outputs[3].is_file():
                raise ValueError('XLSX exporter returned without saving the requested workbook')
        print('\n'.join(text[:3]))
        print('Report:', outputs[3] if runtime else report_path)
        return 0
    except (ValueError,ImportError,OSError,subprocess.CalledProcessError) as exc:
        parser.exit(2,f'ERROR: {exc}\n')


if __name__=='__main__':
    raise SystemExit(main())

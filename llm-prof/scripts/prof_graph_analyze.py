#!/usr/bin/env python3
"""Graph-off source mapping + graph-on device timing for CUDA/HIP/SDAA traces."""

import argparse
import csv
import json
import math
from collections import Counter
from pathlib import Path

from graph_source_mapping import (
    attach_sources, build_mapping, complete, correlation, device_events,
    is_graph_launch, is_runtime,
)
from prof_analyze import analyze_kernels, analyze_kernels_detailed, load_trace
from kernel_categories import add_classification_args, configure_classifier, get_rule_provenance


def window_arg(value):
    try:
        start, end = map(float, value.split(":"))
        if not math.isfinite(start) or not math.isfinite(end) or end <= start:
            raise ValueError()
        return start, end
    except ValueError as exc:
        raise argparse.ArgumentTypeError("Expected START_US:END_US with end > start") from exc


def graph_evidence(events, kernels):
    ids = {correlation(k) for k in kernels} - {None}
    names = Counter(str(e["name"]) for e in events if complete(e) and is_runtime(e)
                    and is_graph_launch(e) and correlation(e) in ids)
    return {"linked_launch_count": sum(names.values()), "linked_launch_names": dict(names)}


def analyze_pair(mapping_events, formal_events, stage, *, mapping_window=None,
                 formal_window=None, mapping_pid=None, formal_pid=None,
                 mapping_device=None, formal_device=None, extra_categories=(), configs=None):
    mapping_kernels = device_events(mapping_events, extra_categories, mapping_pid,
                                    mapping_device, mapping_window)
    formal_kernels = device_events(formal_events, extra_categories, formal_pid,
                                   formal_device, formal_window)
    if not formal_kernels:
        raise ValueError("No formal device events; inspect categories/selection. CPU durations are not a fallback.")
    mapping = build_mapping(mapping_events, mapping_kernels)
    warnings = ["Source attribution is inferred across traces by symbol, not per graph node or replay instance.",
                "Stage is caller-selected; mixed traces require explicit windows."]
    graph_signals = {"mapping": graph_evidence(mapping_events, mapping_kernels),
                     "formal": graph_evidence(formal_events, formal_kernels)}
    if graph_signals["mapping"]["linked_launch_count"]:
        warnings.append("Mapping device events include graph replay links; verify graph-off capture and window.")
    if not graph_signals["formal"]["linked_launch_count"]:
        warnings.append("No graph launch linked to selected formal kernels; graph-on execution must be verified from logs/exporter.")
    differences = []
    if configs is None:
        warnings.append("Run configuration compatibility is unverified; retain both serve configs and compare manually.")
    else:
        if not isinstance(configs, dict) or not all(isinstance(configs.get(k), dict)
                                                   for k in ("mapping", "formal")):
            raise ValueError("Run configs require mapping/formal objects")
        off, on = configs.get("mapping", {}), configs.get("formal", {})
        for key in sorted((set(off) | set(on)) - {"graph_enabled"}):
            if key not in off or key not in on or off[key] != on[key]:
                differences.append({"field": key, "mapping": off.get(key), "formal": on.get(key)})
        if off.get("graph_enabled") is not False or on.get("graph_enabled") is not True:
            raise ValueError("Run configs must declare mapping.graph_enabled=false and formal.graph_enabled=true")
        if differences:
            warnings.append("Run configs differ; source matches are flagged ambiguous.")
    if not mapping_kernels:
        warnings.append("No mapping device events; all source mappings will be unresolved.")
    total = sum(e["dur"] for e in formal_kernels)
    rows = analyze_kernels_detailed(formal_kernels, 0)
    # No absolute percentage/idle from summed durations: kernels may overlap.
    for row in rows:
        row.pop("absolute_pct", None)
    attach_sources(rows, mapping, differences)
    counts = Counter(row["source_status"] for row in rows)
    time_by_status = Counter()
    for row in rows:
        time_by_status[row["source_status"]] += row["total_duration_us"]
    return {"schema_version": 1, "stage": stage,
            "classification_rules": get_rule_provenance(),
            "timing_source": "formal/graph-on/device-events",
            "source_mapping_source": "mapping/graph-off",
            "run_config_status": "unverified" if configs is None else "different" if differences else "declared-equal",
            "duration_unit": "us", "formal_kernel_count": len(formal_kernels),
            "formal_total_duration_us": total,
            "mapping_kernel_count": len(mapping_kernels),
            "graph_evidence": graph_signals,
            "category_summary": analyze_kernels(formal_kernels),
            "source_summary": {"kernel_symbols_by_status": dict(counts),
                               "formal_duration_us_by_status": dict(time_by_status)},
            "pair_config_differences": differences, "warnings": warnings, "kernels": rows}


def write_report(report, output_dir):
    output_dir.mkdir(parents=True, exist_ok=True)
    (output_dir / "graph_analysis.json").write_text(
        json.dumps(report, ensure_ascii=False, indent=2), encoding="utf8")
    columns = ["kernel_name", "category", "classification_rule", "call_count",
               "total_duration_us", "avg_duration_us", "relative_pct", "source_status",
               "match_basis", "ambiguity_reasons", "source_candidates", "semantic_type",
               "classification_confidence", "classification_ambiguity", "classification_rule_id",
               "classification_evidence", "classification_category_candidates", "rules_version", "rules_sha256"]
    records = []
    for row in report["kernels"]:
        row = dict(row, rules_version=report["classification_rules"]["version"],
                   rules_sha256=report["classification_rules"]["sha256"])
        records.append({key: json.dumps(row[key], ensure_ascii=False) if isinstance(row[key], list)
                        else row[key] for key in columns})
    with (output_dir / "graph_kernel_sources.csv").open("w", newline="", encoding="utf-8-sig") as handle:
        writer = csv.DictWriter(handle, fieldnames=columns)
        writer.writeheader()
        writer.writerows(records)
    lines = [f"Stage: {report['stage']}", f"Timing: {report['timing_source']}",
             "Rules: " + json.dumps(report["classification_rules"], ensure_ascii=False),
             f"Device duration sum (us): {report['formal_total_duration_us']:.3f}",
             "Kernel durations are summed, not wall time or utilization.",
             "Source status: " + json.dumps(report["source_summary"], ensure_ascii=False)]
    lines.extend("WARNING: " + warning for warning in report["warnings"])
    for row in report["kernels"]:
        lines.append(f"{row['total_duration_us']:.3f} us | {row['kernel_name']} | "
                     f"{row['source_status']} | {row['match_basis']} | "
                     + ", ".join(row["ambiguity_reasons"]))
    (output_dir / "graph_analysis_summary.txt").write_text("\n".join(lines) + "\n", encoding="utf8")
    try:
        from openpyxl import Workbook
    except ImportError:
        return
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "Graph-on kernel sources"
    sheet.append(columns)
    for record in records:
        # Strings are literal data, including vendor names starting with '='.
        sheet.append([value[:32000] + " [TRUNCATED: see graph_analysis.json]"
                      if isinstance(value, str) and len(value) > 32000 else value
                      for value in (record[key] for key in columns)])
        for cell in sheet[sheet.max_row]:
            if isinstance(cell.value, str):
                cell.data_type = "s"
    sheet.freeze_panes = "A2"
    sheet.auto_filter.ref = sheet.dimensions
    from openpyxl.utils import get_column_letter
    for index in range(1, len(columns) + 1):
        sheet.column_dimensions[get_column_letter(index)].width = 24 if index not in (1, 11, 16) else 65
    categories = workbook.create_sheet("Graph-on categories")
    categories.append(["category", "total_duration_us", "call_count", "relative_pct"])
    for name, stats in report["category_summary"].items():
        categories.append([name, stats["total_dur_ms"] * 1000, stats["count"], stats["percentage"]])
    categories.freeze_panes = "A2"
    for col in "ABCD":
        categories.column_dimensions[col].width = 25
    audit = workbook.create_sheet("Provenance")
    for key, value in report.items():
        if key != "kernels":
            audit.append([key, json.dumps(value, ensure_ascii=False)])
    workbook.save(output_dir / "graph_analysis.xlsx")


def main():
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--mapping-trace", required=True, help="Graph-off JSON(.gz)")
    parser.add_argument("--formal-trace", required=True, help="Graph-on JSON(.gz)")
    parser.add_argument("--stage", required=True, choices=("prefill", "decode", "all"),
                        help="Use stage-separated traces, or provide independent trace windows")
    parser.add_argument("--output-dir", required=True)
    parser.add_argument("--mapping-window", type=window_arg, help="START_US:END_US in mapping trace")
    parser.add_argument("--formal-window", type=window_arg, help="START_US:END_US in formal trace")
    for label in ("mapping", "formal"):
        parser.add_argument(f"--{label}-pid")
        parser.add_argument(f"--{label}-device")
    parser.add_argument("--device-category", action="append", default=[],
                        help="Additional confirmed device execution category (repeatable)")
    parser.add_argument("--run-configs", help="JSON with mapping/formal run metadata and graph_enabled booleans")
    add_classification_args(parser)
    args = parser.parse_args()
    try:
        configure_classifier(args.kernel_rules, args.platform)
        if Path(args.mapping_trace).resolve() == Path(args.formal_trace).resolve():
            raise ValueError("Mapping and formal must be different captures")
        configs = json.loads(Path(args.run_configs).read_text(encoding="utf8")) if args.run_configs else None
        mapping_events = load_trace(args.mapping_trace)
        formal_events = load_trace(args.formal_trace)
        report = analyze_pair(mapping_events, formal_events, args.stage,
                              mapping_window=args.mapping_window, formal_window=args.formal_window,
                              mapping_pid=args.mapping_pid, formal_pid=args.formal_pid,
                              mapping_device=args.mapping_device, formal_device=args.formal_device,
                              extra_categories=args.device_category, configs=configs)
        report["inputs"] = {"mapping_trace": str(Path(args.mapping_trace).resolve()),
                            "formal_trace": str(Path(args.formal_trace).resolve()),
                            "mapping_window_us": args.mapping_window, "formal_window_us": args.formal_window,
                            "mapping_pid": args.mapping_pid, "formal_pid": args.formal_pid,
                            "mapping_device": args.mapping_device, "formal_device": args.formal_device,
                            "extra_device_categories": args.device_category, "run_configs": configs}
        write_report(report, Path(args.output_dir))
    except (ValueError, OSError, UnicodeError) as exc:
        parser.exit(1, f"ERROR: {exc}\n")
    print(f"Graph-on device duration sum: {report['formal_total_duration_us']:.3f} us")
    print("Source status:", report["source_summary"]["kernel_symbols_by_status"])
    for warning in report["warnings"]:
        print("WARNING:", warning)
    print("Reports:", Path(args.output_dir).resolve())


if __name__ == "__main__":
    main()

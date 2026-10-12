#!/usr/bin/env python3
"""Audit kernel classification on trace JSON(.gz) or aggregated XLSX tables.

The audit reports matched rule reasons and unknown fallbacks instead of silently
accepting every fallback as a correct elementwise classification.
"""
from __future__ import annotations

import argparse
import collections
import csv
import gzip
import json
from pathlib import Path

from kernel_categories import (
    add_classification_args, classify_kernel_detail, configure_classifier, get_rule_provenance,
)
from graph_source_mapping import DEVICE_CATEGORIES


def iter_trace_names(path: Path):
    opener = gzip.open if path.suffix == ".gz" else open
    with opener(path, "rt", encoding="utf-8", errors="replace") as f:
        data = json.load(f)
    events = data.get("traceEvents", data.get("events", [])) if isinstance(data, dict) else data
    for event in events:
        if not isinstance(event, dict):
            continue
        if event.get("cat") in DEVICE_CATEGORIES:
            yield str(event.get("name", "")), float(event.get("dur", 0) or 0), 1


def iter_xlsx_names(path: Path):
    try:
        import openpyxl
    except ImportError as exc:
        raise SystemExit("XLSX audit requires openpyxl: pip install openpyxl") from exc
    wb = openpyxl.load_workbook(path, read_only=True, data_only=True)
    try:
        ws = wb.active
        # Stream rows once: repeated ws.cell() re-reads read-only worksheets.
        rows = ws.iter_rows(values_only=True)
        headers = list(next(rows, ()))
        name_header = next((h for h in ("kernel_name", "name") if h in headers), None)
        name_col = headers.index(name_header) if name_header else 0
        duration_col = headers.index("total_duration_us") if "total_duration_us" in headers else None
        count_header = next((h for h in ("call_count", "record_count") if h in headers), None)
        count_col = headers.index(count_header) if count_header else None
        for row in rows:
            name = row[name_col] if len(row) > name_col else None
            if not name:
                continue
            duration = float(row[duration_col] or 0) if duration_col is not None else 0
            count = int(row[count_col] or 1) if count_col is not None else 1
            yield str(name), duration, count
    finally:
        wb.close()


def audit(paths):
    agg = collections.defaultdict(lambda: {"category": None, "reason": None, "duration_us": 0.0, "count": 0})
    for path in paths:
        if path.suffix.lower() == ".xlsx":
            iterator = iter_xlsx_names(path)
        elif path.suffix.lower() in (".tsv", ".csv"):
            iterator = iter_table_names(path)
        else:
            iterator = iter_trace_names(path)
        for name, duration, count in iterator:
            detail = classify_kernel_detail(name)
            category, reason = detail["category"], detail["reason"]
            key = (name, category, reason)
            agg[key]["category"] = category
            agg[key]["reason"] = reason
            for field in ("rule_id", "rule_platform", "semantic_type", "confidence",
                          "evidence", "ambiguity_reasons", "matched_pattern", "category_candidates"):
                agg[key][field] = detail.get(field)
            agg[key]["duration_us"] += duration
            agg[key]["count"] += count
    rows = []
    for (name, category, reason), values in agg.items():
        rows.append({"name": name, **values})
    rows.sort(key=lambda x: x["duration_us"], reverse=True)
    return rows


def iter_table_names(path):
    with path.open(encoding="utf-8-sig", newline="") as handle:
        reader = csv.DictReader(handle, delimiter="\t" if path.suffix.lower() == ".tsv" else ",")
        for row in reader:
            name = row.get("kernel_name", row.get("name"))
            if not name:
                continue
            duration = float(row.get("total_duration_us", row.get("duration_us", 0)) or 0)
            count = int(row.get("call_count", row.get("record_count", row.get("count", 1))) or 1)
            yield name, duration, count


def main():
    parser = argparse.ArgumentParser()
    parser.add_argument("paths", nargs="+", type=Path)
    parser.add_argument("--output", type=Path)
    parser.add_argument("--fail-on-fallback", action="store_true")
    add_classification_args(parser)
    args = parser.parse_args()
    try:
        configure_classifier(args.kernel_rules, args.platform)
    except (ValueError, OSError) as exc:
        parser.error(str(exc))
    rows = audit(args.paths)
    summary = collections.defaultdict(lambda: {"count": 0, "duration_us": 0.0})
    for row in rows:
        key = (row["category"], row["reason"])
        summary[key]["count"] += row["count"]
        summary[key]["duration_us"] += row["duration_us"]
    print("category\treason\tcount\tduration_ms")
    for (category, reason), values in sorted(summary.items(), key=lambda x: -x[1]["duration_us"]):
        print(f"{category}\t{reason}\t{values['count']}\t{values['duration_us'] / 1000:.3f}")
    unknown = [r for r in rows if r["matched_pattern"] is None or r["confidence"] == "unknown"]
    print(f"\nunknown_fallback_names={len(unknown)}")
    for row in unknown[:50]:
        print(f"  {row['duration_us'] / 1000:.3f} ms\t{row['count']} calls\t{row['name'][:180]}")
    if args.output:
        args.output.write_text(json.dumps({"classification_rules": get_rule_provenance(),
                                          "summary": summary_to_json(summary), "rows": rows},
                                         ensure_ascii=False, indent=2), encoding="utf-8")
    if args.fail_on_fallback and unknown:
        raise SystemExit(2)


def summary_to_json(summary):
    return {f"{category}|{reason}": values for (category, reason), values in summary.items()}


if __name__ == "__main__":
    main()

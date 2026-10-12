#!/usr/bin/env python3
"""Attribution ambiguity and graph-on timing invariants across trace exporters."""

import gzip
import json
import subprocess
import sys
import tempfile
import unittest
from pathlib import Path

from graph_source_mapping import IntervalIndex, device_events
from prof_graph_analyze import analyze_pair


def event(name, cat, ts, dur, args=None, pid=42, tid=7):
    return dict(name=name, cat=cat, ph="X", ts=ts, dur=dur, pid=pid, tid=tid, args=args or {})


def capture(name="kernel_123", runtime="cuda_runtime", ext="External id", corr="correlation",
            cpu_name="aten::mm", location="model.py(100): forward", offset=0):
    return [event(location, "python_function", offset, 80),
            event(cpu_name, "cpu_op", offset + 10, 20, {ext: 8}),
            event("vendorLaunchKernel", runtime, offset + 15, 2, {ext: 8, corr: 9}),
            event(name, "kernel", offset + 40, 1000, {ext: 8, corr: 9, "device": 0}, pid=0, tid=19)]


def formal(name="kernel_123", duration=5):
    return [event("cudaGraphLaunch", "cuda_runtime", 0, 99, {"correlation": 500}),
            event("aten::gemm_host_should_not_count", "cpu_op", 0, 9999),
            event(name, "kernel", 100, duration, {"correlation": 500, "device": 0}, pid=0, tid=19),
            event("WAIT", "kernel", 200, 77, {"device": 0}, pid=0, tid=19)]


class GraphMappingTests(unittest.TestCase):
    def test_exporters_and_timing_conservation(self):
        for runtime, ext, corr, name in [
            ("cuda_runtime", "External id", "correlation", "cublasLt_gemm"),
            ("hip_runtime", "external_id", "correlation_id", "Cijk_Alik_Bljk_F8BS_MT16x16x64"),
            ("privateuse1_runtime", "External id", "correlation", "teco_slave_gemm_NN(Args)"),
            ("sdaa_runtime", "externalId", "correlationId", "tecolmk_slave_fused_norm<2,1>(Args)"),
        ]:
            with self.subTest(runtime=runtime):
                report = analyze_pair(capture(name, runtime, ext, corr), formal(name), "decode")
                row = report["kernels"][0]
                self.assertEqual(report["formal_total_duration_us"], 5)
                self.assertEqual(report["formal_kernel_count"], 1)
                self.assertEqual(report["graph_evidence"]["formal"]["linked_launch_count"], 1)
                self.assertEqual(report["graph_evidence"]["mapping"]["linked_launch_count"], 0)
                self.assertEqual(report["category_summary"]["总计"]["total_dur_ms"], .005)
                self.assertEqual(sum(report["source_summary"]["formal_duration_us_by_status"].values()), 5)
                self.assertEqual(row["source_status"], "mapped")
                self.assertEqual(row["match_basis"], "name:exact")
                self.assertEqual(row["source_candidates"][0]["cpu_op"], "aten::mm")
                self.assertNotIn("absolute_pct", row)

    def test_correlation_without_kernel_external_id(self):
        mapping = capture()
        mapping[-1]["args"].pop("External id")
        row = analyze_pair(mapping, formal(), "decode")["kernels"][0]
        self.assertEqual(row["source_status"], "mapped")
        self.assertEqual(row["source_candidates"][0]["attribution_basis"], ["correlation+launch-external-id"])

    def test_host_enclosure_and_python_only(self):
        mapping = capture()
        for e in mapping:
            e["args"].pop("External id", None)
        row = analyze_pair(mapping, formal(), "decode")["kernels"][0]
        self.assertEqual(row["source_candidates"][0]["attribution_basis"], ["correlation+host-thread-enclosure"])
        mapping = [e for e in mapping if e["cat"] != "cpu_op"]
        row = analyze_pair(mapping, formal(), "decode")["kernels"][0]
        self.assertEqual(row["source_status"], "partial")
        self.assertIsNone(row["source_candidates"][0]["cpu_op"])

    def test_multi_layer_ambiguity_and_missing_occurrence(self):
        mapping = capture() + capture(location="model.py(200): other_layer", offset=2000)
        # Give the second observation unique trace-local IDs.
        for e in mapping[4:]:
            if "External id" in e["args"]:
                e["args"]["External id"] = 18
            if "correlation" in e["args"]:
                e["args"]["correlation"] = 19
        row = analyze_pair(mapping, formal(), "decode")["kernels"][0]
        self.assertEqual(row["source_status"], "ambiguous")
        self.assertEqual(len(row["source_candidates"]), 2)
        self.assertIn("multiple-source-sites", row["ambiguity_reasons"])
        mapping = capture() + [event("kernel_123", "kernel", 2000, 1000, {"device": 0}, pid=0)]
        row = analyze_pair(mapping, formal(), "decode")["kernels"][0]
        self.assertIn("mapping-occurrences-without-source", row["ambiguity_reasons"])

    def test_normalization_preserves_template_and_numeric_suffixes(self):
        row = analyze_pair(capture("void teco_gemm<32>(Params)"), formal("teco_gemm<32>"), "decode")["kernels"][0]
        self.assertEqual(row["match_basis"], "name:normalized")
        for name in ("teco_gemm<64>", "teco_gemm<32>_v2"):
            row = analyze_pair(capture("void teco_gemm<32>(Params)"), formal(name), "decode")["kernels"][0]
            self.assertEqual(row["source_status"], "unresolved")

    def test_multiple_normalized_symbols_not_picked_arbitrarily(self):
        mapping = capture("void kernel_123(int)") + capture("kernel_123(float)", offset=2000)
        row = analyze_pair(mapping, formal(), "decode")["kernels"][0]
        self.assertIn("multiple-mapping-symbols", row["ambiguity_reasons"])
        self.assertEqual(row["source_status"], "ambiguous")

    def test_graph_replay_scope_not_original_source(self):
        mapping = capture()
        mapping[2]["name"] = "sdaaGraphLaunch"
        row = analyze_pair(mapping, formal(), "decode")["kernels"][0]
        self.assertEqual(row["source_status"], "unresolved")
        self.assertIn("graph-launch-only", row["ambiguity_reasons"])

    def test_correlation_reused_across_threads_and_graphs(self):
        mapping = capture()
        mapping.append(event("sdaaGraphLaunch", "privateuse1_runtime", 20, 1,
                             {"correlation": 9}, tid=8))
        row = analyze_pair(mapping, formal(), "decode")["kernels"][0]
        self.assertEqual(row["source_status"], "ambiguous")
        self.assertIn("mixed-graph-and-kernel-launch-links", row["ambiguity_reasons"])
        mapping[-1]["tid"] = 7
        row = analyze_pair(mapping, formal(), "decode")["kernels"][0]
        self.assertEqual(row["source_status"], "unresolved")

    def test_config_schema_and_unverified_status(self):
        self.assertEqual(analyze_pair(capture(), formal(), "decode")["run_config_status"], "unverified")
        configs = {"mapping": {"graph_enabled": False}, "formal": {"graph_enabled": True}}
        self.assertEqual(analyze_pair(capture(), formal(), "decode", configs=configs)["run_config_status"], "declared-equal")
        for bad in ([], {}, {"mapping": [], "formal": {}}):
            with self.assertRaisesRegex(ValueError, "mapping/formal"):
                analyze_pair(capture(), formal(), "decode", configs=bad)

    def test_cpu_only_mapping_and_cpu_only_formal(self):
        mapping = [event("aten::gemm", "cpu_op", 0, 20)]
        report = analyze_pair(mapping, formal(), "decode")
        self.assertEqual(report["kernels"][0]["source_status"], "unresolved")
        with self.assertRaisesRegex(ValueError, "No formal device"):
            analyze_pair(capture(), mapping, "decode")
        with self.assertRaisesRegex(ValueError, "host/runtime"):
            device_events(mapping, ["cpu_op"])

    def test_multi_device_is_explicit_and_windows_independent(self):
        events = formal() + [event("another", "kernel", 120, 10, {"device": 1}, pid=1)]
        with self.assertRaisesRegex(ValueError, "Multiple device"):
            analyze_pair(capture(), events, "decode")
        report = analyze_pair(capture(), events, "decode", formal_device=0,
                              mapping_window=(30, 1100), formal_window=(100, 110))
        self.assertEqual(report["formal_total_duration_us"], 5)

    def test_config_differences_flagged(self):
        configs = {"mapping": {"graph_enabled": False, "batch": 1},
                   "formal": {"graph_enabled": True, "batch": 2}}
        row = analyze_pair(capture(), formal(), "decode", configs=configs)["kernels"][0]
        self.assertIn("pair-config-differences", row["ambiguity_reasons"])
        self.assertEqual(row["source_status"], "ambiguous")
        configs["formal"]["graph_enabled"] = False
        with self.assertRaisesRegex(ValueError, "graph_enabled"):
            analyze_pair(capture(), formal(), "decode", configs=configs)

    def test_nested_interval_index_matches_enclosure(self):
        events = [event("root", "python_function", 0, 10000)]
        events += [event(str(i), "python_function", i, 2) for i in range(1, 1000)]
        index = IntervalIndex(events)
        for timestamp in (0, 15.5, 999.5, 1000, 9999, 10000):
            self.assertEqual({e["name"] for e in index.at(timestamp)},
                             {e["name"] for e in events if e["ts"] <= timestamp < e["ts"] + e["dur"]})

    def test_real_cli_gzip_reports_and_provenance(self):
        with tempfile.TemporaryDirectory() as directory:
            root = Path(directory)
            for label, events in [("mapping", capture(runtime="privateuse1_runtime")), ("formal", formal())]:
                with gzip.open(root / f"{label}.json.gz", "wt", encoding="utf8") as handle:
                    json.dump({"traceEvents": events}, handle)
            command = [sys.executable, str(Path(__file__).with_name("prof_graph_analyze.py")),
                       "--mapping-trace", str(root / "mapping.json.gz"),
                       "--formal-trace", str(root / "formal.json.gz"),
                       "--stage", "decode", "--output-dir", str(root / "report")]
            result = subprocess.run(command, capture_output=True, text=True)
            self.assertEqual(result.returncode, 0, result.stderr)
            report = json.loads((root / "report/graph_analysis.json").read_text(encoding="utf8"))
            self.assertEqual(report["formal_total_duration_us"], 5)
            self.assertEqual(report["kernels"][0]["source_status"], "mapped")
            self.assertEqual(report["inputs"]["formal_trace"], str((root / "formal.json.gz").resolve()))
            self.assertTrue((root / "report/graph_kernel_sources.csv").exists())
            self.assertTrue((root / "report/graph_analysis_summary.txt").exists())
            try:
                from openpyxl import load_workbook
            except ImportError:
                return
            workbook = load_workbook(root / "report/graph_analysis.xlsx")
            self.assertEqual(workbook.active.cell(2, 5).value, 5)
            self.assertEqual(workbook.active.cell(2, 8).value, "mapped")
            self.assertIn("Provenance", workbook.sheetnames)
            self.assertIn("Graph-on categories", workbook.sheetnames)

    def test_graph_xlsx_category_audit_preserves_calls(self):
        try:
            import openpyxl
        except ImportError:
            self.skipTest("openpyxl not installed")
        from audit_kernel_categories import audit
        from prof_graph_analyze import write_report
        events = formal() + [event("kernel_123", "kernel", 300, 7, {"device": 0}, pid=0)]
        report = analyze_pair(capture(), events, "decode")
        with tempfile.TemporaryDirectory() as directory:
            write_report(report, Path(directory))
            rows = audit([Path(directory) / "graph_analysis.xlsx"])
            self.assertEqual(rows[0]["count"], 2)
            self.assertEqual(rows[0]["duration_us"], 12)

    def test_multistream_overlap_is_not_bubble(self):
        events = formal() + [event("second_kernel", "kernel", 101, 8,
                                  {"device": 0, "stream": 20}, pid=0, tid=20)]
        report = analyze_pair(capture(), events, "decode")
        self.assertEqual(report["formal_total_duration_us"], 13)
        self.assertEqual(sum(r["total_duration_us"] for r in report["kernels"]), 13)
        self.assertNotIn("bubble", report)

    def test_vendor_device_category_and_invalid_windows(self):
        events = formal()
        events[2]["cat"] = "vendor_device_execution"
        report = analyze_pair(capture(), events, "decode", extra_categories=["vendor_device_execution"])
        self.assertEqual(report["formal_total_duration_us"], 5)
        from prof_graph_analyze import window_arg
        import argparse
        for invalid in ("nan:1", "0:inf", "1:1", "bad"):
            with self.assertRaises(argparse.ArgumentTypeError):
                window_arg(invalid)


if __name__ == "__main__":
    unittest.main()
